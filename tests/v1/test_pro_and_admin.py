"""Pro entitlement and administrator access, enforced by the API (the web app's hidden buttons are not the
check). Throwaway database with the plans the migration seeds; the workbook itself is built from the real NMB
report when it is loaded locally (skipped in CI, which has no DSE data)."""
from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import apps.api.main as api
from packages.database.models import Entitlement, Plan, User
from pipelines import accounts
from tests.v1.test_auth_and_portfolios import _signup, db  # noqa: F401  (fixture)

SHEETS = ["Executive Summary", "Company Profile", "Income Statement", "Balance Sheet", "Cash Flow",
          "Financial Ratios", "Bank Metrics", "Valuation", "Valuation Assumptions", "Sensitivity Analysis",
          "Technical Analysis", "Risk Analysis", "Sources & Evidence"]


@pytest.fixture()
def client(db):  # noqa: F811
    with db() as s:
        s.add_all([Plan(id="free", name="Free", description="", is_active=True),
                   Plan(id="pro", name="Pro", description="", is_active=True)])
        s.flush()
        s.add(Entitlement(plan_id="pro", feature="excel_export"))
        s.commit()
    return TestClient(api.app)


def _set(db, email, **fields):  # noqa: F811
    with db() as s:
        u = s.query(User).filter_by(email=email).one()
        for k, v in fields.items():
            setattr(u, k, v)
        s.commit()


def test_excel_export_needs_a_session(client):
    r = client.get("/api/v1/reports/DSE:TEST/xlsx")
    assert r.status_code == 401


def test_a_free_account_is_refused_by_the_api_itself(client):
    h = _signup(client, "free@example.com")
    r = client.get("/api/v1/reports/DSE:TEST/xlsx", headers=h)
    assert r.status_code == 403 and "Pro" in r.json()["detail"]
    me = client.get("/api/v1/auth/me", headers=h).json()
    assert me["plan"]["id"] == "free" and me["features"] == [] and me["role"] == "user"


def test_a_pro_account_gets_past_the_entitlement(client, db):  # noqa: F811
    h = _signup(client, "pro@example.com")
    _set(db, "pro@example.com", plan_id="pro")
    assert client.get("/api/v1/auth/me", headers=h).json()["features"] == ["excel_export"]
    # DSE:TEST has no annual reports, so past the entitlement the answer is the report's own 404.
    assert client.get("/api/v1/reports/DSE:TEST/xlsx", headers=h).status_code == 404


def test_admin_data_health_is_for_administrators_only(client, db):  # noqa: F811
    h = _signup(client, "someone@example.com")
    assert client.get("/api/v1/admin/data-health").status_code == 401
    assert client.get("/api/v1/admin/data-health", headers=h).status_code == 403
    _set(db, "someone@example.com", role="admin")
    r = client.get("/api/v1/admin/data-health", headers=h)
    assert r.status_code == 200
    body = r.json()
    assert {p["id"] for p in body["providers"]} == {"dse_public", "mansa", "eodhd"}
    mansa = next(p for p in body["providers"] if p["id"] == "mansa")
    # The setting is named so the administrator knows what to set; its value never appears.
    assert mansa["credential_env"] == "MANSA_API_KEY" and mansa["state"] in ("NOT_CONFIGURED", "LICENSE_REVIEW_REQUIRED")


def test_the_accounts_command_grants_a_plan_and_a_role(client, db, monkeypatch, capsys):  # noqa: F811
    _signup(client, "owner@example.com")
    monkeypatch.setattr(accounts, "SessionLocal", db)
    assert accounts.main(["grant-plan", "--email", "Owner@Example.com", "--plan", "pro"]) == 0
    assert accounts.main(["set-role", "--email", "owner@example.com", "--role", "admin"]) == 0
    assert "plan pro, role admin, features ['excel_export']" in capsys.readouterr().out
    assert accounts.main(["grant-plan", "--email", "owner@example.com", "--plan", "platinum"]) == 1
    assert accounts.main(["show", "--email", "nobody@example.com"]) == 1


# ------------------------------------------------------------------ the workbook (real NMB data)

def _nmb_report():
    try:
        from packages.database.session import SessionLocal
        from packages.report.builder import build_report
        with SessionLocal() as s:
            return build_report(s, "DSE:NMB")
    except Exception:
        return None


_REPORT = _nmb_report()
needs_nmb = pytest.mark.skipif(_REPORT is None, reason="NMB data not loaded locally")


@needs_nmb
def test_the_workbook_has_thirteen_sheets_and_copies_the_report_exactly():
    from packages.report.excel import build_workbook
    wb = load_workbook(io.BytesIO(build_workbook(_REPORT)))
    assert wb.sheetnames == SHEETS
    # The latest net interest income in the Income Statement sheet is the report's figure, to the shilling.
    ws = wb["Income Statement"]
    years = [str(y) for y in _REPORT["years"]]
    header = [c.value for c in ws[4]]
    col = header.index(f"FY{years[-1]}") + 1
    row = next(r for r in ws.iter_rows(min_row=5) if r[0].value == "Net interest income")
    want = next(x for x in _REPORT["statements"][0]["rows"] if x["item_code"] == "net_interest_income")
    assert row[col - 1].value == want["cells"][years[-1]]["value"]
    # The fair value on the summary is the valuation's, and the sources sheet lists every statement figure.
    summary = {r[0].value: r[1].value for r in wb["Executive Summary"].iter_rows(min_row=5) if r[0].value}
    # Excel keeps about 15 significant digits, so a computed value matches to that precision, not to 28.
    fv = _REPORT["valuation"]["result"]["fair_value"]
    assert abs(summary["Fair value (per share)"] - float(fv)) < 1e-9 * float(fv)
    figures = sum(1 for st in _REPORT["statements"] for r in st["rows"] for y in years if r["cells"].get(y))
    evidence = [r for r in wb["Sources & Evidence"].iter_rows(min_row=1) if r[0].value in
                {st["title"] for st in _REPORT["statements"]}]
    assert len(evidence) == figures


@needs_nmb
def test_an_unavailable_figure_is_written_as_its_status_never_as_a_number():
    from packages.report.excel import build_workbook
    wb = load_workbook(io.BytesIO(build_workbook(_REPORT)))
    words = {c.value for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
    gaps = [st for st in _REPORT["statements"] for r in st["rows"] for c in r["cells"].values() if not c.get("available")]
    if gaps:
        assert words & {"INSUFFICIENT DATA", "CONFLICTING SOURCE", "BLOCKED"}


@needs_nmb
def test_a_pro_account_downloads_the_workbook(client, db, monkeypatch):  # noqa: F811
    h = _signup(client, "analyst@example.com")
    _set(db, "analyst@example.com", plan_id="pro")
    monkeypatch.setattr(api, "get_report", lambda security_id, session: _REPORT)
    r = client.get("/api/v1/reports/DSE:NMB/xlsx", headers=h)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    assert 'filename="AfriEdge_DSE_NMB_' in r.headers["content-disposition"]
    assert load_workbook(io.BytesIO(r.content)).sheetnames == SHEETS
